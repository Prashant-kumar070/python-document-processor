from io import BytesIO

from docx import Document as DocxDocument
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app.core.config import settings
from app.db.database import Base, get_db
from app.main import app
from app.core.security import get_current_user
from app.services.llm_service import LLMProviderError


def test_upload_to_chat_returns_actual_stored_source(monkeypatch, tmp_path):
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(bind=engine)
    upload_dir = tmp_path / "uploads"
    monkeypatch.setattr(settings, "UPLOAD_DIR", str(upload_dir))
    monkeypatch.setattr(settings, "DATABASE_URL", "sqlite://")
    monkeypatch.setattr(settings, "EMBEDDING_PROVIDER", "mock")
    monkeypatch.setattr(settings, "MIN_SIMILARITY_SCORE", -1.0)

    def override_get_db():
        with Session(engine) as session:
            yield session

    app.dependency_overrides[get_db] = override_get_db
    llm = RecordingLLM()
    monkeypatch.setattr("app.services.rag_service.LLMService", lambda: llm)

    document = DocxDocument()
    document.add_paragraph("JWT authentication is described below.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Feature"
    table.cell(0, 1).text = "Detail"
    table.cell(1, 0).text = "JWT"
    table.cell(1, 1).text = "Signed bearer tokens verify each request."
    document.add_paragraph("End of document.")
    file_content = BytesIO()
    document.save(file_content)
    file_content.seek(0)

    try:
        with TestClient(app) as client:
            registration = client.post(
                "/auth/register",
                json={
                    "name": "Pipeline user",
                    "email": "pipeline-user@example.com",
                    "password": "test-password-123",
                },
            )
            assert registration.status_code == 201
            token = client.post(
                "/auth/login",
                json={
                    "email": "pipeline-user@example.com",
                    "password": "test-password-123",
                },
            ).json()["access_token"]
            headers = {"Authorization": f"Bearer {token}"}

            upload = client.post(
                "/documents",
                headers=headers,
                files={
                    "file": (
                        "authentication.docx",
                        file_content.getvalue(),
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    )
                },
            )
            assert upload.status_code == 201
            document_id = upload.json()["id"]

            detail = client.get(f"/documents/{document_id}", headers=headers)
            assert detail.status_code == 200
            assert detail.json()["status"] == "COMPLETED"
            assert detail.json()["chunk_count"] == 1

            chat = client.post(
                "/chat",
                headers=headers,
                json={
                    "question": "How does JWT authentication verify a request?",
                    "document_ids": [document_id],
                },
            )
            assert chat.status_code == 200
            response = chat.json()
            assert response["answer"] == "The document says JWT uses a signed bearer token [Source 1]."
            assert len(response["sources"]) == 1
            assert response["sources"][0]["document_id"] == document_id
            assert response["sources"][0]["filename"] == "authentication.docx"
            assert response["sources"][0]["page_number"] == 1
            assert response["sources"][0]["chunk_id"] > 0

            conversation = client.get(
                f"/conversations/{response['conversation_id']}",
                headers=headers,
            )
            assert conversation.status_code == 200
            assistant_message = conversation.json()["messages"][-1]
            assert assistant_message["role"] == "assistant"
            assert assistant_message["sources"][0]["document_chunk_id"] == response["sources"][0]["chunk_id"]
    finally:
        app.dependency_overrides.pop(get_db, None)
        Base.metadata.drop_all(bind=engine)
        engine.dispose()


class RecordingLLM:
    def generate_answer(self, *, system_instruction, messages):
        assert "| Feature | Detail |" in messages[-1]["content"]
        assert "| JWT | Signed bearer tokens verify each request. |" in messages[-1]["content"]
        assert "End of document." in messages[-1]["content"]
        assert "using ONLY" in system_instruction
        return "The document says JWT uses a signed bearer token [Source 1]."


def test_chat_endpoint_returns_structured_upstream_provider_error(monkeypatch):
    def override_current_user():
        return type("CurrentUser", (), {"id": 1})()

    def override_database():
        yield object()

    class FailingRagService:
        def __init__(self, db):
            pass

        def answer_question(self, **kwargs):
            raise LLMProviderError(
                "Quota exceeded for this model.",
                provider="Gemini",
                provider_status=429,
                provider_code="RESOURCE_EXHAUSTED",
            )

    app.dependency_overrides[get_current_user] = override_current_user
    app.dependency_overrides[get_db] = override_database
    monkeypatch.setattr("app.api.routes.chat.RagService", FailingRagService)
    try:
        with TestClient(app) as client:
            response = client.post(
                "/chat",
                headers={"Authorization": "Bearer test-token"},
                json={"question": "Question"},
            )
    finally:
        app.dependency_overrides.pop(get_current_user, None)
        app.dependency_overrides.pop(get_db, None)

    assert response.status_code == 503
    assert response.json()["detail"] == {
        "error": "LLM provider request failed",
        "provider": "Gemini",
        "message": "Quota exceeded for this model.",
        "provider_status": 429,
        "provider_code": "RESOURCE_EXHAUSTED",
    }
