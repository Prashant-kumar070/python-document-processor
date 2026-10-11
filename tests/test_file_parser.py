from docx import Document

from app.utils.file_parser import parse_docx


def test_parse_docx_includes_paragraphs_and_table_rows_in_document_order(tmp_path):
    document = Document()
    document.add_paragraph("Document introduction.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Requirement"
    table.cell(0, 1).text = "Description"
    table.cell(1, 0).text = "Authentication"
    table.cell(1, 1).text = "Signed bearer tokens validate requests."
    document.add_paragraph("Document conclusion.")
    file_path = tmp_path / "with-table.docx"
    document.save(file_path)

    pages = parse_docx(str(file_path))

    assert len(pages) == 1
    text = pages[0]["text"]
    assert text.index("Document introduction.") < text.index("| Requirement | Description |")
    assert text.index("| Authentication | Signed bearer tokens") < text.index(
        "Document conclusion."
    )
    assert "Signed bearer tokens validate requests." in text


def test_parse_docx_extracts_nested_tables_inside_cell(tmp_path):
    document = Document()
    outer_table = document.add_table(rows=1, cols=1)
    cell = outer_table.cell(0, 0)
    cell.text = "Outer cell"
    nested_table = cell.add_table(rows=1, cols=2)
    nested_table.cell(0, 0).text = "Nested key"
    nested_table.cell(0, 1).text = "Nested value"
    file_path = tmp_path / "nested-table.docx"
    document.save(file_path)

    text = parse_docx(str(file_path))[0]["text"]

    assert "Outer cell" in text
    assert "| Nested key | Nested value |" in text


def test_parse_docx_includes_headers_body_tables_and_footers(tmp_path):
    document = Document()
    section = document.sections[0]
    section.header.paragraphs[0].text = "Company Confidential"
    section.footer.paragraphs[0].text = "Page footer text"
    document.add_paragraph("Main body text.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Table label"
    table.cell(0, 1).text = "Table value"
    file_path = tmp_path / "header-body-footer.docx"
    document.save(file_path)

    text = parse_docx(str(file_path))[0]["text"]

    assert "[HEADER]\nCompany Confidential" in text
    assert "Main body text." in text
    assert "| Table label | Table value |" in text
    assert "[FOOTER]\nPage footer text" in text
    assert text.index("[HEADER]") < text.index("Main body text.")
    assert text.index("Table value") < text.index("[FOOTER]")
